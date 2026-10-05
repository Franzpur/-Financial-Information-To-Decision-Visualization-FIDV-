#!/usr/bin/env python3
"""Map DATA-SPACE BICS member workbooks onto L4 hierarchy codes.

Sources: 20261003 L1 workbooks (with Level4 columns) and 20261005 Staples
one-file-per-L4 workbooks (L4 from filename). Skips the old L1-only Staples
book, pull-* formula sheets, Office locks, and the formula example.
Does not commit xlsx. Output: bics_entities_20261003.db
"""
from __future__ import annotations

import json
import sqlite3
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("need openpyxl: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
HIER_JSON = HERE / "bics-equity-hierarchy-2024.json"
SRC_DIRS = (
    REPO / "DATA-SPACE" / "ICBC C" / "20261003" / "BICS_LEGALENTITY",
    REPO / "DATA-SPACE" / "ICBC C" / "20261005",
)
DB_OUT = HERE / "bics_entities_20261003.db"

L1_FROM_FILE = {
    "BICS_Communications.xlsx": "Communications",
    "BICS_Comsumer Discretionary.xlsx": "Consumer Discretionary",
    "BICS_Energy.xlsx": "Energy",
    "BICS_Financials.xlsx": "Financials",
    "BICS_Health Care.xlsx": "Health Care",
    "BICS_Industrials.xlsx": "Industrials",
    "BICS_Materials.xlsx": "Materials",
    "BICS_Real Estate.xlsx": "Real Estate",
    "BICS_Technology.xlsx": "Technology",
    "BICS_Utilities.xlsx": "Utilities",
    "__BICS_Government__.xlsx": "Government",
    "BICS_Government.xlsx": "Government",
}

STAPLES_L1 = "Consumer Staples"
FILENAME_L4_ALIASES = {
    "agricultural & producers": "agricultural producers",
}


def cellstr(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def load_hierarchy() -> tuple[dict[tuple[str, int], dict], dict[str, dict], dict[str, dict]]:
    data = json.loads(HIER_JSON.read_text(encoding="utf-8"))
    by_name: dict[tuple[str, int], dict] = {}
    by_code: dict[str, dict] = {}
    staples_l4: dict[str, dict] = {}
    for n in data.get("nodes", []):
        name = (n.get("name") or "").strip()
        lvl = n.get("level")
        code = str(n.get("bicsCode") or "")
        if name and lvl:
            by_name[(name.lower(), int(lvl))] = n
        if code:
            by_code[code] = n
        if int(lvl or 0) == 4 and code.startswith("12"):
            staples_l4[norm_key(name)] = n
    return by_name, by_code, staples_l4


def norm_key(s: str) -> str:
    t = (s or "").lower().replace("&", " ").replace("-", " ")
    return " ".join(t.split())


def l4_from_filename(path: Path, staples_l4: dict[str, dict]) -> dict | None:
    stem = path.stem.strip()
    alias = FILENAME_L4_ALIASES.get(stem.lower(), stem.lower())
    return staples_l4.get(norm_key(alias)) or staples_l4.get(norm_key(stem))


def ancestors(node: dict, by_code: dict[str, dict]) -> tuple[str, str, str]:
    l1 = l2 = l3 = ""
    cur: dict | None = node
    seen = 0
    while cur and seen < 8:
        seen += 1
        lvl = cur.get("level")
        name = (cur.get("name") or "").strip()
        if lvl == 1:
            l1 = name
        elif lvl == 2:
            l2 = name
        elif lvl == 3:
            l3 = name
        pc = cur.get("parentCode")
        cur = by_code.get(pc) if pc else None
    return l1, l2, l3


def header_map(row) -> dict[str, int]:
    out = {}
    for i, c in enumerate(row):
        h = cellstr(c).lower()
        if h:
            out[h] = i
    return out


def col(hmap: dict[str, int], *needles: str) -> int | None:
    for needle in needles:
        for h, i in hmap.items():
            if needle in h:
                return i
    return None


def as_float(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if s in ("#N/A", "n/a", "N/A"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def should_skip(path: Path) -> bool:
    n = path.name.lower()
    if n.startswith("~$"):
        return True
    if "fomular" in n:
        return True
    if n.startswith("pull-"):
        return True
    if "comsumer staples" in n or n.startswith("bics_consumer staples"):
        return True
    return False


def _at(row, idx) -> str:
    if idx is None or idx >= len(row):
        return ""
    return cellstr(row[idx])


def _ingest_rows(
    wb,
    path: Path,
    by_name: dict,
    by_code: dict,
    staples_l4: dict,
    cur: sqlite3.Cursor,
) -> tuple[int, int]:
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        return 0, 0
    hmap = header_map(header)
    i_co = col(hmap, "member companies", "company")
    i_tk = col(hmap, "member ticker", "ticker")
    i_mkt = col(hmap, "mkt cap")
    i_rev = col(hmap, "ind rev")
    i_pct = col(hmap, "% tot rev")
    i_l2 = col(hmap, "level2", "level 2")
    i_l3 = col(hmap, "level3", "level 3")
    i_l4 = col(hmap, "level4", "level 4")
    file_node = l4_from_filename(path, staples_l4) if i_l4 is None else None
    if i_l4 is None and file_node is None:
        return 0, 0
    ok = 0
    reject = 0
    file_l1 = L1_FROM_FILE.get(path.name, "")
    for r in rows:
        if not any(c not in (None, "") for c in r):
            continue
        node = None
        l4 = ""
        if i_l4 is not None:
            l4 = _at(r, i_l4)
            if l4 and l4 not in ("#N/A", "n/a"):
                node = by_name.get((l4.lower(), 4))
        else:
            node = file_node
            l4 = (node.get("name") or "") if node else ""
        if not node:
            reject += 1
            continue
        if file_l1:
            l1 = file_l1
            l2 = _at(r, i_l2)
            l3 = _at(r, i_l3)
        else:
            l1, l2, l3 = ancestors(node, by_code)
            if not l1:
                l1 = STAPLES_L1
        cur.execute(
            """INSERT INTO entity_memberships
               (name, ticker, l1_name, l2_name, l3_name, l4_name, bics_code_l4,
                legal_entity_coord, mkt_cap, ind_rev, pct_tot_rev, source_file)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                _at(r, i_co),
                _at(r, i_tk),
                l1,
                l2,
                l3,
                l4,
                node.get("bicsCode"),
                node.get("legalEntityCoord"),
                as_float(r[i_mkt]) if i_mkt is not None and i_mkt < len(r) else None,
                as_float(r[i_rev]) if i_rev is not None and i_rev < len(r) else None,
                as_float(r[i_pct]) if i_pct is not None and i_pct < len(r) else None,
                path.name,
            ),
        )
        ok += 1
    return ok, reject


def ingest_file(
    path: Path,
    by_name: dict,
    by_code: dict,
    staples_l4: dict,
    cur: sqlite3.Cursor,
) -> tuple[int, int]:
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        return _ingest_rows(wb, path, by_name, by_code, staples_l4, cur)
    finally:
        wb.close()


def source_files() -> list[Path]:
    files: list[Path] = []
    for folder in SRC_DIRS:
        if not folder.is_dir():
            print(f"missing source dir {folder}", file=sys.stderr)
            sys.exit(1)
        files.extend(p for p in folder.glob("*.xlsx") if not should_skip(p))
    files.sort(key=lambda p: (str(p.parent), p.name.lower()))
    return files


def main() -> None:
    if not HIER_JSON.is_file():
        print("missing hierarchy JSON", HIER_JSON, file=sys.stderr)
        sys.exit(1)
    by_name, by_code, staples_l4 = load_hierarchy()
    if DB_OUT.exists():
        DB_OUT.unlink()
    conn = sqlite3.connect(DB_OUT)
    cur = conn.cursor()
    cur.execute(
        """CREATE TABLE entity_memberships (
            id INTEGER PRIMARY KEY,
            name TEXT,
            ticker TEXT,
            l1_name TEXT,
            l2_name TEXT,
            l3_name TEXT,
            l4_name TEXT,
            bics_code_l4 TEXT NOT NULL,
            legal_entity_coord TEXT,
            mkt_cap REAL,
            ind_rev REAL,
            pct_tot_rev REAL,
            source_file TEXT
        )"""
    )
    cur.execute(
        "CREATE INDEX idx_entities_l4 ON entity_memberships(bics_code_l4)"
    )
    cur.execute(
        "CREATE INDEX idx_entities_ticker ON entity_memberships(ticker)"
    )
    total_ok = 0
    total_rej = 0
    for path in source_files():
        ok, rej = ingest_file(path, by_name, by_code, staples_l4, cur)
        print(f"{path.parent.name}/{path.name}: kept={ok} reject={rej}")
        total_ok += ok
        total_rej += rej
    conn.commit()
    n_codes = cur.execute(
        "SELECT COUNT(DISTINCT bics_code_l4) FROM entity_memberships"
    ).fetchone()[0]
    conn.close()
    print(f"wrote {DB_OUT} rows={total_ok} reject={total_rej} l4_codes={n_codes}")


if __name__ == "__main__":
    main()
