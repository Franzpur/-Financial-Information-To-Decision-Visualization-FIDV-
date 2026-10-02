#!/usr/bin/env python3
"""Build C-COORD-3 industry coordinate library from Bloomberg BICS Equity Hierarchy 2024.

Source: 2144414.xlsx sheet "BICS Equity Hierarchy 2024"
Outputs (same directory):
  - bics-equity-hierarchy-2024.json  (canonical)
  - bics_hierarchy.db                (SQLite mirror)

Re-run anytime; idempotent.
"""
from __future__ import annotations

import json
import sqlite3
import sys
from datetime import datetime, timezone
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("need openpyxl: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "2144414.xlsx"
SHEET = "BICS Equity Hierarchy 2024"
JSON_OUT = HERE / "bics-equity-hierarchy-2024.json"
DB_OUT = HERE / "bics_hierarchy.db"


def compact_code(raw) -> str:
    if raw is None:
        return ""
    if isinstance(raw, float):
        raw = int(raw)
    s = str(raw).strip()
    if s.endswith(".0") and s.replace(".", "", 1).isdigit():
        s = s[:-2]
    return s


def to_legal_entity_coord(bics_code: str) -> str:
    """7 segments × 2 digits, hyphenated; pad right with 00 for shallow leaves."""
    digits = "".join(c for c in bics_code if c.isdigit())
    if len(digits) % 2:
        raise ValueError(f"odd-length BICS code: {bics_code!r}")
    segs = [digits[i : i + 2] for i in range(0, len(digits), 2)]
    if len(segs) > 7:
        raise ValueError(f"more than 7 levels: {bics_code!r}")
    while len(segs) < 7:
        segs.append("00")
    return "-".join(segs)


def parent_of(bics_code: str) -> str | None:
    if len(bics_code) <= 2:
        return None
    return bics_code[:-2]


def _cell(row, idx: int, default=""):
    if len(row) <= idx or row[idx] is None:
        return default
    return row[idx]


def _parse_hierarchy_row(row) -> dict | None:
    code = compact_code(_cell(row, 2, None))
    level_raw = _cell(row, 3, None)
    if not code or level_raw is None:
        return None
    try:
        level = int(level_raw)
    except (TypeError, ValueError):
        return None
    return {
        "bicsCode": code,
        "level": level,
        "name": str(_cell(row, 4) or "").strip(),
        "nameLong": str(_cell(row, 5) or "").strip(),
        "definition": str(_cell(row, 6) or "").strip(),
        "nameZh": str(_cell(row, 15) or "").strip(),
    }


def load_rows(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        wb.close()
        raise SystemExit(f"sheet not found: {SHEET}")
    ws = wb[SHEET]
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i < 5:
            continue
        parsed = _parse_hierarchy_row(row)
        if parsed:
            rows.append(parsed)
    wb.close()
    return rows


def build_nodes(rows: list[dict]) -> list[dict]:
    codes = {r["bicsCode"] for r in rows}
    # A node is a leaf if no other code is a longer proper prefix extension (+2 digits at a time)
    has_child = set()
    for c in codes:
        p = parent_of(c)
        if p and p in codes:
            has_child.add(p)

    nodes = []
    for r in rows:
        code = r["bicsCode"]
        level = r["level"]
        if len(code) != 2 * level:
            raise ValueError(f"code length != 2*level: {code} level={level}")
        parent = parent_of(code)
        if level > 1 and (parent is None or parent not in codes):
            raise ValueError(f"missing parent for {code} (want {parent})")
        is_leaf = code not in has_child
        nodes.append(
            {
                "bicsCode": code,
                "level": level,
                "leafLevel": level if is_leaf else None,
                "name": r["name"],
                "nameLong": r["nameLong"],
                "definition": r["definition"],
                "nameZh": r["nameZh"],
                "parentCode": parent,
                "isLeaf": is_leaf,
                "legalEntityCoord": to_legal_entity_coord(code),
            }
        )
    nodes.sort(key=lambda n: (n["level"], n["bicsCode"]))
    return nodes


def write_json(nodes: list[dict]) -> None:
    payload = {
        "version": "BICS Equity Hierarchy 2024",
        "concept": "C-COORD-3",
        "sourceFile": SOURCE.name,
        "sheet": SHEET,
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "codeFormat": {
            "authoritative": "bicsCode compact digits, length 2*level",
            "productField": "legalEntityCoord = 7×2 hyphenated, right-pad 00",
            "example": "10-10-10-10-12-10-00",
        },
        "stats": {
            "nodes": len(nodes),
            "leaves": sum(1 for n in nodes if n["isLeaf"]),
            "byLevel": {str(L): sum(1 for n in nodes if n["level"] == L) for L in range(1, 8)},
            "leafDepth": {
                str(L): sum(1 for n in nodes if n["isLeaf"] and n["level"] == L)
                for L in range(1, 8)
            },
        },
        "nodes": nodes,
    }
    JSON_OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("wrote", JSON_OUT, "nodes=", len(nodes), "leaves=", payload["stats"]["leaves"])
    print("byLevel", payload["stats"]["byLevel"])
    print("leafDepth", payload["stats"]["leafDepth"])


def write_sqlite(nodes: list[dict]) -> None:
    if DB_OUT.exists():
        DB_OUT.unlink()
    con = sqlite3.connect(DB_OUT)
    con.execute(
        """
        CREATE TABLE bics_nodes (
          bics_code TEXT PRIMARY KEY,
          level INTEGER NOT NULL,
          name TEXT NOT NULL,
          name_long TEXT,
          definition TEXT,
          name_zh TEXT,
          parent_code TEXT,
          is_leaf INTEGER NOT NULL,
          legal_entity_coord TEXT NOT NULL
        )
        """
    )
    con.executemany(
        """
        INSERT INTO bics_nodes (
          bics_code, level, name, name_long, definition, name_zh,
          parent_code, is_leaf, legal_entity_coord
        ) VALUES (?,?,?,?,?,?,?,?,?)
        """,
        [
            (
                n["bicsCode"],
                n["level"],
                n["name"],
                n["nameLong"],
                n["definition"],
                n["nameZh"],
                n["parentCode"],
                1 if n["isLeaf"] else 0,
                n["legalEntityCoord"],
            )
            for n in nodes
        ],
    )
    con.execute("CREATE INDEX idx_bics_parent ON bics_nodes(parent_code)")
    con.execute("CREATE INDEX idx_bics_level ON bics_nodes(level)")
    con.execute("CREATE INDEX idx_bics_coord ON bics_nodes(legal_entity_coord)")
    con.commit()
    con.close()
    print("wrote", DB_OUT)


def main() -> None:
    if not SOURCE.exists():
        raise SystemExit(f"missing source: {SOURCE}")
    rows = load_rows(SOURCE)
    nodes = build_nodes(rows)
    write_json(nodes)
    write_sqlite(nodes)


if __name__ == "__main__":
    main()
